# export_controls.py
import json, re, sys, argparse
from pathlib import Path
from typing import Any, List, Dict, Iterable, Tuple

# Excel
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_CONTROLS_ROOT = BASE_DIR / "controls"


# ---------- helpers ----------
def _loads_any(text: str) -> Any:
    """Load JSON that might itself contain a JSON string."""
    data = json.loads(text)
    if isinstance(data, str):
        data = json.loads(data)
    return data


def _as_list(obj: Any) -> List[dict]:
    if isinstance(obj, list):
        return obj
    if isinstance(obj, dict):
        return [obj]
    raise ValueError("Expected a JSON object or array.")


def _parse_recommendations(val: Any) -> List[str]:
    """
    Normalize recommendations to a numbered list of steps.
    """
    if isinstance(val, list):
        # Always number list items
        return [f"{i+1}) {str(x).strip()}" for i, x in enumerate(val) if str(x).strip()]
    if isinstance(val, str):
        s = val.strip().replace("\\n", "\n")
        try:
            inner = json.loads(s)
            if isinstance(inner, list):
                return [f"{i+1}) {str(x).strip()}" for i, x in enumerate(inner) if str(x).strip()]
        except Exception:
            pass
        # Split numbered text and keep numbering
        parts = re.split(r"(?:^|\n)\s*\d+\)\s*", s)
        steps = [p.strip() for p in parts if p.strip()]
        if steps:
            # return [f"\n{i+1}) {step}" for i, step in enumerate(steps)]
            return [f"{i+1}) {step}" for i, step in enumerate(steps)]
        return [f"1) {s}"] if s else []
    return []


def _normalize_controls(data: Iterable[dict]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for item in data:
        rec_steps = _parse_recommendations(
            item.get("recommendations") or item.get("recommendation_details")
        )
        rec_text = "\r\n".join(rec_steps)  # <-- real Excel line breaks

        out.append({
            "application":     item.get("application", ""),
            "url":             item.get("url", ""),
            "control_subject": item.get("control_subject") or item.get("name", ""),
            "description":     item.get("description", ""),
            "category":        item.get("category", ""),
            "severity":        item.get("severity", ""),
            "recommendations": rec_text,      # <-- use processed text
            "additional_info": item.get("additional_info", ""),
        })
    return out


# ---------- writers ----------
def save_pretty_json(data: List[dict], output_prefix: Path) -> Path:
    pretty_path = output_prefix.with_name(output_prefix.name + "_pretty.json")
    pretty_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return pretty_path


def save_xlsx(rows: List[Dict[str, Any]], output_prefix: Path) -> Path:
    xlsx_path = output_prefix.with_suffix(".xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "controls"

    headers = [
        "application", "url", "control_subject", "description",
        "category", "severity", "recommendations", "additional_info"
    ]
    ws.append(headers)

    # --- WRITE DATA ROWS ---
    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    # wrap text (esp. needed for recommendations newlines)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                  min_col=1, max_col=len(headers)):
        for cell in row_cells:
            cell.alignment = wrap

    # autosize columns a bit
    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        base_width = 15
        if header in ("url", "description", "recommendations"):
            base_width = 60
        elif header == "control_subject":
            base_width = 35
        max_len = max(len(str(ws.cell(r, col_idx).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[col_letter].width = min(max(base_width, int(max_len * 0.9)), 80)

    # (optional) set row heights so all steps are visible
    rec_col = headers.index("recommendations") + 1
    for r in range(2, ws.max_row + 1):
        txt = str(ws.cell(row=r, column=rec_col).value or "")
        # '\r\n' becomes '\n' when read back, so count '\n'
        lines = txt.count("\n") + 1
        ws.row_dimensions[r].height = min(18 * lines, 180)

    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    return xlsx_path


def main_convert_controls(input_path: str | Path, output_prefix: str | Path | None = None) -> tuple[Path, Path]:
    input_path = Path(input_path)
    raw = input_path.read_text(encoding="utf-8")

    data = _loads_any(raw)
    controls = _as_list(data)
    rows = _normalize_controls(controls)

    if output_prefix is None:
        output_prefix = input_path.with_suffix("")
    output_prefix = Path(output_prefix)

    # ensure output folder exists
    output_prefix.parent.mkdir(parents=True, exist_ok=True)

    pretty = save_pretty_json(controls, output_prefix)
    xlsx = save_xlsx(rows, output_prefix)
    return pretty, xlsx


def convert_file(input_path: str | Path, output_prefix: str | Path | None = None) -> Tuple[Path, Path]:
    """Convert by explicit file path."""
    input_path = Path(input_path)
    if not input_path.exists():
        raise FileNotFoundError(f"Input JSON not found: {input_path}\nCWD={Path.cwd()}")
    return main_convert_controls(input_path, output_prefix)


def convert_controls(app_name: str, base_dir: Path = DEFAULT_CONTROLS_ROOT) -> Tuple[Path, Path]:
    """
    Wrapper by app name.
    Looks in:
      1) controls/<app>/<app>_controls.json
      2) controls/<app>/<app>.json
      3) controls/<app>_controls.json (root fallback)
      4) controls/<app>.json        (root fallback)

    Writes to:
      controls/<app>/<app>_pretty.json
      controls/<app>/<app>.xlsx
    """

    # In case someone passes a path, keep only the last segment (e.g., "Dropbox")
    app_key = Path(app_name).name.strip()

    controls_root = Path(base_dir)  # already absolute from DEFAULT_CONTROLS_ROOT
    folder = controls_root / app_key

    # preferred locations in subfolder
    candidates = [
        folder / f"{app_key}_controls.json",
        folder / f"{app_key}.json",
    ]
    
    # root fallbacks (if you saved JSON at controls/<app>_controls.json)
    root_fallbacks = [
        controls_root / f"{app_key}_controls.json",
        controls_root / f"{app_key}.json",
    ]

    input_path = next((p for p in candidates if p.exists()), None)

    if input_path is None:
        input_path = next((p for p in root_fallbacks if p.exists()), None)

    if input_path is None:
        
        found_sub = [p.name for p in folder.glob("*.json")] if folder.exists() else []
        found_root = [p.name for p in controls_root.glob("*.json")]
        raise FileNotFoundError()

    # outputs always go into the subfolder
    folder.mkdir(parents=True, exist_ok=True)
    output_prefix = folder / app_key

    return main_convert_controls(input_path, output_prefix)
